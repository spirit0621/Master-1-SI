import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_inline_formatted(paragraph, text, base_font_name=None, base_size=None, base_color=None):
    # Regex to extract bold, italic, inline code, links
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        if base_font_name:
            run.font.name = base_font_name
        if base_size:
            run.font.size = base_size
        if base_color:
            run.font.color.rgb = base_color

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                run.text = m.group(1)
                run.font.color.rgb = RGBColor(0, 102, 204)
                run.underline = True
            else:
                run.text = token
        else:
            run.text = token

def convert_markdown_to_docx(md_path, docx_path):
    if not os.path.exists(md_path):
        raise FileNotFoundError(f"Fichier source introuvable: {md_path}")
        
    doc = docx.Document()
    
    # Configuration des marges de page (2 cm partout)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles par défaut
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    
    def flush_code_block():
        nonlocal in_code_block, code_lines
        if code_lines:
            code_text = "".join(code_lines).rstrip()
            # Créer un bloc tableau 1x1 pour encadrer le code avec fond gris clair
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.8)
            set_cell_background(cell, "F4F6F9")
            set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(30, 30, 30)
            
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_lines = []
        in_code_block = False

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
            
        parsed_rows = []
        for line in table_lines:
            line_str = line.strip()
            if not line_str.startswith('|'):
                continue
            cells = [c.strip() for c in line_str.split('|')[1:-1]]
            # Ignorer la ligne de séparation |:---|---:|
            if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                continue
            parsed_rows.append(cells)
            
        if parsed_rows:
            col_count = max(len(r) for r in parsed_rows)
            # Égaliser le nombre de colonnes
            normalized_rows = [r + [''] * (col_count - len(r)) for r in parsed_rows]
            
            tbl = doc.add_table(rows=len(normalized_rows), cols=col_count)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = True
            
            for row_idx, row_data in enumerate(normalized_rows):
                row = tbl.rows[row_idx]
                is_header = (row_idx == 0)
                for col_idx, cell_value in enumerate(row_data):
                    cell = row.cells[col_idx]
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    if is_header:
                        set_cell_background(cell, "003366") # Bleu marine corporate
                    else:
                        if row_idx % 2 == 1:
                            set_cell_background(cell, "FFFFFF")
                        else:
                            set_cell_background(cell, "F7F9FC") # Ligne alternée gris bleuté
                            
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    
                    if is_header:
                        add_inline_formatted(p, cell_value, base_font_name='Calibri', base_size=Pt(10), base_color=RGBColor(255, 255, 255))
                        for r in p.runs:
                            r.bold = True
                    else:
                        add_inline_formatted(p, cell_value, base_font_name='Calibri', base_size=Pt(9.5), base_color=RGBColor(40, 40, 40))
                        
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        table_lines = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Gestion des blocs de code ```
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Gestion des tableaux markdown
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                flush_table()
                
        # Ligne vide
        if not stripped:
            i += 1
            continue
            
        # Ligne horizontale ---
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 45)
            run.font.color.rgb = RGBColor(180, 190, 200)
            run.bold = True
            i += 1
            continue
            
        # Titre H1 #
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = doc.add_heading(level=1)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            add_inline_formatted(p, h_text, base_font_name='Calibri', base_size=Pt(20), base_color=RGBColor(0, 51, 102))
            i += 1
            continue
            
        # Titre H2 ##
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_heading(level=2)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            add_inline_formatted(p, h_text, base_font_name='Calibri', base_size=Pt(15), base_color=RGBColor(0, 80, 150))
            i += 1
            continue

        # Titre H3 ###
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_heading(level=3)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            add_inline_formatted(p, h_text, base_font_name='Calibri', base_size=Pt(12.5), base_color=RGBColor(50, 100, 160))
            i += 1
            continue

        # Titre H4 ####
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_heading(level=4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            add_inline_formatted(p, h_text, base_font_name='Calibri', base_size=Pt(11), base_color=RGBColor(70, 70, 70))
            i += 1
            continue
            
        # Citation / Callout >
        if stripped.startswith('> '):
            callout_text = stripped[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.8)
            set_cell_background(cell, "EEF4FA") # Bleu très doux
            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline_formatted(p, callout_text, base_font_name='Calibri', base_size=Pt(10), base_color=RGBColor(0, 51, 102))
            for r in p.runs:
                r.italic = True
            i += 1
            continue
            
        # Liste à puces (- ou *)
        if re.match(r'^[-*]\s+', stripped):
            item_text = re.sub(r'^[-*]\s+', '', stripped)
            # Gestion checkbox [x] ou [ ]
            if item_text.startswith('[x] ') or item_text.startswith('[X] '):
                item_text = "☑  " + item_text[4:]
            elif item_text.startswith('[ ] '):
                item_text = "☐  " + item_text[4:]
                
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted(p, item_text)
            i += 1
            continue
            
        # Liste numérotée (1. 2. etc.)
        if re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted(p, item_text)
            i += 1
            continue
            
        # Paragraphe standard
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_formatted(p, stripped)
        i += 1
        
    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()
        
    doc.save(docx_path)
    return docx_path

if __name__ == '__main__':
    base_docs = r"c:\Users\alves\Desktop\Lycée, bts , formation, master\CFA-insta\Master 1 SI\TP\TPFINALE\docs"
    base_tp = r"c:\Users\alves\Desktop\Lycée, bts , formation, master\CFA-insta\Master 1 SI\TP\TPFINALE"
    
    files = [
        (os.path.join(base_docs, "RAPPORT_TECHNIQUE_TECHCORP.md"), os.path.join(base_docs, "RAPPORT_TECHNIQUE_TECHCORP.docx")),
        (os.path.join(base_docs, "LIVRABLES_ET_CONTENU_DU_RAPPORT.md"), os.path.join(base_docs, "LIVRABLES_ET_CONTENU_DU_RAPPORT.docx")),
        (os.path.join(base_tp, "README.md"), os.path.join(base_docs, "README.docx"))
    ]
    
    for src, dst in files:
        res = convert_markdown_to_docx(src, dst)
        print(f"Généré : {res} ({os.path.getsize(res)} octets)")
